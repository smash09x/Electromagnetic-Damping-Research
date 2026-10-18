"""Reproduce every number, table and figure in the paper from the raw DAQ dumps.

Usage:  python run_all.py
Reads   Data/Raw/
Writes  Data/Results/  (final_table.csv, final_stats.json, clock_calibration.json,
        decay_halves.json, video_analysis.json, supporting_numbers.json, Overall_Results.xlsx, figures/, logs/)
The video frame times in Data/Video/frame_times/ come from the original .mp4 files via
check_video_clocks.py (needs the videos and ffprobe; not part of this run).
"""
import csv, json, os, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.abspath(os.path.join(HERE, "..", "..", "Data", "Results"))
LOGS = os.path.join(RES, "logs")
os.makedirs(LOGS, exist_ok=True)

STEPS = [
    ("clock_calibration.py",      "Clock factor from the 50 Hz mains line"),
    ("alternative_estimators.py", "Peak, per-cycle rms and NLLS decay estimators"),
    ("transduction_model.py",     "Physical transduction-model fit"),
    ("final_table.py",            "Headline estimator, clock-corrected table (Table A1)"),
    ("significance_tests.py",     "1/R fit, significance and 2.2 MOhm divider check"),
    ("video_analysis.py",         "Video tracks: independent clock check and decay of the motion"),
    ("decay_halves.py",           "Early- vs late-half decay rates"),
    ("supporting_numbers.py",     "ADC step, noise, turns, air drag, time per read"),
    ("figures.py",                "Figures"),
]

for i, (script, what) in enumerate(STEPS, 1):
    print(f"[{i}/{len(STEPS)}] {what}")
    r = subprocess.run([sys.executable, script], cwd=HERE, capture_output=True, text=True)
    with open(os.path.join(LOGS, script.replace(".py", ".txt")), "w") as fh:
        fh.write(r.stdout + r.stderr)
    if r.returncode:
        sys.exit(f"{script} failed; see {LOGS}")

# One workbook with everything, like Overall_Results.xlsx in the companion repository
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = "Per-run results"
with open(os.path.join(RES, "final_table.csv")) as fh:
    for row in csv.reader(fh):
        ws.append([float(x) if x.replace(".", "", 1).replace("-", "", 1).isdigit() else x for x in row])
st = json.load(open(os.path.join(RES, "final_stats.json")))
ws2 = wb.create_sheet("Summary")
ws2.append(["Quantity", "Spring 1", "Spring 2"])
for k, lab in (("gamma", "gamma (1e-3 /s), mean"), ("Q_mean", "Q, mean"), ("f_true_mean", "f (Hz)")):
    a, b = st["S1"][k], st["S2"][k]
    ws2.append([lab, a[0] if isinstance(a, list) else a, b[0] if isinstance(b, list) else b])
c = st["S1_cfit"]
ws2.append([]); ws2.append(["1/R fit (spring 1): c measured (Ohm/s)", c["n8_lt10k"]["c"], c["n8_lt10k"]["c_err"]])
ws2.append(["c predicted (Ohm/s)", c["c_pred"]])
ws2.append(["2.2 MOhm divider: amplitude-corrected ratio", st["divider_2M2"]["ratio_amp_corrected"], "predicted", st["divider_2M2"]["predicted"]])
clk = json.load(open(os.path.join(RES, "clock_calibration.json")))
ws2.append(["Clock factor s (mains hum): value, run sd, grid tolerance", clk["s"], clk["s_sd"], clk["s_grid_tol"]])
vid = json.load(open(os.path.join(RES, "video_analysis.json")))
for sp in (1, 2):
    o = vid[f"spring{sp}"]
    ws2.append([f"Clock factor s (video, spring {sp})", o["s_video"], o["s_video_se"]])
o = vid["spring1"]
ws2.append(["Decay of the motion, spring 1 video (1e-3 /s): mean, sd", o["gamma_video_mean"], o["gamma_video_sd"]])
ws3 = wb.create_sheet("Video tracks")
keys = ["video", "spring", "fps", "frame_dt_sd_ms", "f_even_spacing", "f", "gamma", "gamma_se", "early", "late", "A0_cm", "resid_mm", "passes_tracking_cut"]
ws3.append(keys)
for r in vid["videos"]:
    ws3.append([r[k] for k in keys])
wb.save(os.path.join(RES, "Overall_Results.xlsx"))
print("done:", RES)
